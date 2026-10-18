"""Гаралт: НЭГ файл; компани бүр өөрийн нэрээр sheet-тэй. Sheet дотор тайлан бүр нэг багана (толгойд тайлангийн огноо), мөр нь бүх компанид ижил."""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import COMPANIES
from .models import ParsedReport, Period
from .ratios import FIRST_ROW_NOTE, GROUPS, Ctx
from .template import ITEMS, TEMPLATE, Section

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
TOTAL_FONT = Font(bold=True)
NUMBER_FORMAT = "#,##0;[Red]-#,##0"
SIDE = Side(style="medium", color="1F3864")
FIRST_DATA_ROW = 10

HEADER_ROWS = (  # (мөрийн дугаар, нэр)
    (4, "Тайлангийн огноо"),
    (5, "Үе"),
    (6, "Хамрах хугацаа (сар)"),
    (7, "Ирүүлсэн огноо (МХБ)"),
    (8, "Эх файл"),
    (9, "Засвар, тэмдэглэл"),
)
_BAD_SHEET_CHARS = set('[]:*?/\\')


def sheet_title(ticker: str, used: set[str]) -> str:
    """Компанийн нэр = sheet нэр (Excel: ≤31 тэмдэгт, [ ] : * ? / \\ хориотой, давхцахгүй)."""
    name = "".join(ch for ch in COMPANIES.get(ticker, ticker) if ch not in _BAD_SHEET_CHARS).strip()[:31] or ticker
    if name.lower() in used:
        name = f"{name[:26]} {ticker}"[:31]
    used.add(name.lower())
    return name


def _visible_rows(reports: list[ParsedReport]):
    """Ямар ч тайланд утга байхгүй мөр, мөн хоосон гарчгийг оруулахгүй (бүх багана ижил мөртэй байна)."""
    used = {k for r in reports for k in r.values}
    rows: list = []
    pending: list[Section] = []     # ихэнх гарчиг нь дараагийн гарчигтай залгаа тул (баланс -> эргэлтийн хөрөнгө) жагсаалт
    for entry in TEMPLATE:
        if isinstance(entry, Section):
            # Шинэ гарчиг эхлэхэд өмнөх хоосон, адил/доод түвшний гарчгийг хаяна (дээд түвшнийг хадгална)
            pending = [p for p in pending if p.level < entry.level]
            pending.append(entry)
        elif entry.key in used:
            rows += pending
            pending = []
            rows.append(entry)
    return rows


def _write_company(ws, ticker: str, cols: list[ParsedReport], rows: list) -> None:
    ws["A1"] = f"{COMPANIES.get(ticker, ticker)} ({ticker}) — санхүүгийн тайлан"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Нэгж: мянган төгрөг. Баланс — эцсийн үлдэгдэл; орлогын тайлан — оны эхнээс тайлангийн огноо хүртэл хуримтлагдсан."
    ws["A2"].font = Font(italic=True, color="595959")
    ws.column_dimensions["A"].width = 78

    if not cols:
        ws["A4"] = "Excel хэлбэрийн тайлан олдсонгүй (зөвхөн PDF байж болно) — mse_reports/ хавтас болон логийг үзнэ үү."
        ws["A4"].font = Font(italic=True, color="C00000")
        return

    for row, name in HEADER_ROWS:
        ws.cell(row=row, column=1, value=name).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=row, column=1).fill = HEADER_FILL

    last = len(cols) + 1
    for i, rep in enumerate(cols, start=2):
        values = {
            4: rep.end_date,
            5: rep.period.label,
            6: rep.months,
            7: date.fromisoformat(rep.filed) if rep.filed else None,
            8: rep.source,
            9: rep.scale_note,
        }
        for row, v in values.items():
            cell = ws.cell(row=row, column=i, value=v)
            cell.fill = HEADER_FILL
            cell.font = Font(bold=True, color="FFFFFF", size=8 if row in (8, 9) else 11)
            cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
            if row in (4, 7):
                cell.number_format = "yyyy-mm-dd"
        ws.column_dimensions[get_column_letter(i)].width = 17

    r = FIRST_DATA_ROW
    row_of: dict[str, int] = {}
    for entry in rows:
        ws.cell(row=r, column=1, value=entry.label)
        if isinstance(entry, Section):
            for c in range(1, last + 1):
                ws.cell(row=r, column=c).fill = SECTION_FILL
                ws.cell(row=r, column=c).font = Font(bold=True)
        else:
            row_of[entry.key] = r
            for i, rep in enumerate(cols, start=2):
                v = rep.values.get(entry.key)
                if v is not None:
                    cell = ws.cell(row=r, column=i, value=round(v, 2))
                    cell.number_format = NUMBER_FORMAT
            if entry.kind == "total":
                for c in range(1, last + 1):
                    ws.cell(row=r, column=c).font = TOTAL_FONT
        r += 1

    # Шалгалт: Нийт хөрөнгө = Өр төлбөр + Эздийн өмч (зарим тайланд "Өр төлбөр ба эздийн өмчийн дүн" мөр байдаггүй)
    if {"total_assets", "total_liabilities", "total_equity"} <= row_of.keys():
        r += 1
        ws.cell(row=r, column=1, value="Балансын зөрүү (хөрөнгө − өр төлбөр − эздийн өмч, 0 байх ёстой)").font = Font(italic=True)
        for i in range(2, last + 1):
            col = get_column_letter(i)
            ta, tl, eq = (f"{col}{row_of[k]}" for k in ("total_assets", "total_liabilities", "total_equity"))
            cell = ws.cell(row=r, column=i,
                           value=f'=IF(AND(ISNUMBER({ta}),ISNUMBER({tl}),ISNUMBER({eq})),ROUND({ta}-{tl}-{eq},0),"")')
            cell.number_format = NUMBER_FORMAT
            cell.font = Font(italic=True)

    _write_ratios(ws, r + 2, row_of, last)

    ws.row_dimensions[8].height = 32
    ws.row_dimensions[9].height = 30
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=2)
    if not any(rep.scale_note for rep in cols):
        ws.row_dimensions[9].hidden = True      # засвар хийгдээгүй бол мөрийг нууна


def _write_ratios(ws, start: int, row_of: dict[str, int], last: int) -> None:
    """Харьцааны шинжилгээний блок (Excel томьёо). Эхлээд мөр бүрийн дугаарыг олгоод дараа нь томьёог бичнэ."""
    last_col = get_column_letter(last)
    r = start
    ws.cell(row=r, column=1, value="ХАРЬЦААНЫ ШИНЖИЛГЭЭ (дээрх мөрүүдээс томьёогоор тооцогдоно)")
    for c in range(1, last + 1):
        ws.cell(row=r, column=c).fill = SECTION_FILL
        ws.cell(row=r, column=c).font = Font(bold=True)
    r += 1
    ws.cell(row=r, column=1, value=FIRST_ROW_NOTE).font = Font(italic=True, color="595959", size=9)
    r += 1
    plan: list[tuple] = []          # (мөр, бүлгийн гарчиг эсэх, ...)
    calc: dict[str, int] = {}
    for title, specs in GROUPS:
        plan.append((r, title, None))
        r += 1
        for spec in specs:
            calc[spec[0]] = r
            plan.append((r, None, spec))
            r += 1
        r += 1
    for row, title, spec in plan:
        if title:
            cell = ws.cell(row=row, column=1, value=title)
            cell.font = Font(bold=True, color="1F3864")
            continue
        sid, label, fmt, build = spec
        ws.cell(row=row, column=1, value=label)
        for i in range(2, last + 1):
            formula = build(Ctx(get_column_letter(i), last_col, row_of, calc))
            if formula:
                cell = ws.cell(row=row, column=i, value=formula)
                cell.number_format = fmt


def build_workbook(reports: list[ParsedReport], tickers: list[str], out_path: Path) -> Path:
    """Нэг файл: тикер бүрт нэг sheet (компанийн нэрээр), тайлан бүр баганаар (огноогоор эрэмбэлсэн)."""
    rows = _visible_rows(reports)               # бүх sheet ижил мөртэй
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for ticker in tickers:
        cols = sorted((r for r in reports if r.ticker == ticker), key=lambda r: r.end_date)
        ws = wb.create_sheet(sheet_title(ticker, used))
        _write_company(ws, ticker, cols, rows)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


_KEY_BY_LABEL = {item.label: item.key for item in ITEMS}


def _to_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def load_existing(path: Path, tickers: list[str]) -> list[ParsedReport]:
    """Өмнө бүтээсэн Excel-ээс тайлан бүрийг (багана бүрийг) буцааж уншина. Sheet бүрийн толгойн мөр (4–9) ба
    мөрийн нэрээр стандарт загварын түлхүүрийг сэргээнэ. Танигдахгүй мөр/sheet-ийг алгасна."""
    wb = load_workbook(path, data_only=False)
    used: set[str] = set()
    out: list[ParsedReport] = []
    for ticker in tickers:
        title = sheet_title(ticker, used)
        if title not in wb.sheetnames:
            continue
        ws = wb[title]
        rows = {}
        for r in range(FIRST_DATA_ROW, ws.max_row + 1):
            key = _KEY_BY_LABEL.get(ws.cell(row=r, column=1).value)
            if key:
                rows[key] = r
        for c in range(2, ws.max_column + 1):
            end = _to_date(ws.cell(row=4, column=c).value)
            m = re.fullmatch(r"(\d{4}) Q([1-4])", str(ws.cell(row=5, column=c).value or ""))
            if end is None or not m:
                continue
            values = {k: float(v) for k, r in rows.items()
                      if isinstance(v := ws.cell(row=r, column=c).value, (int, float))}
            filed = _to_date(ws.cell(row=7, column=c).value)
            out.append(ParsedReport(
                ticker, Period(int(m[1]), int(m[2])), end, int(ws.cell(row=6, column=c).value or int(m[2]) * 3),
                values=values, source=str(ws.cell(row=8, column=c).value or ""),
                filed=filed.isoformat() if filed else "", scale_note=str(ws.cell(row=9, column=c).value or "")))
    wb.close()
    return out
