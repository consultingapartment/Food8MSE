"""Excel (.xls/.xlsx) тайлангаас Баланс ба Орлого/ашиг алдагдлын хүснэгтийг стандарт загварт задлах."""
from __future__ import annotations

import io

import logging
import re
from datetime import date
from pathlib import Path
from typing import Any, Optional

from .models import Downloaded, ParsedReport, Period
from .english import english_key, is_english
from .template import BS, BY_KEY, GROUP_HEADERS, IGNORED, IS, ITEMS, group_of, match_item, norm

log = logging.getLogger(__name__)

BALANCE_SHEET_RE = re.compile(r"^(bs|бс|сбд|сбт|баланс|нсбт|nst|ст)\b|баланс|санхүү(гийн)?\s*байдал|санхүү\s*байдл", re.I)
INCOME_SHEET_RE = re.compile(r"^(is|одт|от|ot|нот)\b|орлого|орлогын|ашиг|үр\s*дүн", re.I)
LABEL_HEADER_RE = re.compile(r"^(үзүүлэлт|балансынзүйл|items?|description|indicators?)")
CODE_HEADER_RE = re.compile(r"^(№|мөрийн\s*дугаар|дугаар)")
NOTE_HEADER_RE = re.compile(r"^(тайлбар|тэмдэглэл|хавсралт|note)")      # утга биш багана (сүүлийн багана гэж андуурахгүй)
DATE_RANGE_RE = re.compile(r"(\d{4})[./-](\d{1,2})[./-](\d{1,2})\s*[-–—]\s*(\d{4})[./-](\d{1,2})[./-](\d{1,2})")
# "2026-01-01 -с 2026-06-30 хүртэл" (Монос хүнсний шинэ маягт)
DATE_RANGE_RE2 = re.compile(r"(\d{4})-(\d{1,2})-(\d{1,2})\s*-?\s*с\s+(\d{4})-(\d{1,2})-(\d{1,2})")
# "2026 оны 6-р сарын 30" (Эм Жи Эл Акуагийн маягт: баганын толгой ба гарчиг дээр дан огноо)
DATE_LONG_RE = re.compile(r"(\d{4})\s*оны\s*(\d{1,2})\s*-?\s*р?\s*сарын\s*(\d{1,2})")
EMBEDDED_CODE_RE = re.compile(r"^(\d+(?:\.\d+)*)(?:\.\s*|\s+)(\S.*)$")        # "1.1.3 Дансны авлага" (дугаар нь нэрний нүдэнд)
TOTAL_NUM_RE = re.compile(r"^нийт\s+\d+(?:\.\d+)*\s+", re.I)          # "Нийт 1.1 Эргэлтийн хөрөнгө" -> "Нийт Эргэлтийн хөрөнгө"
#: Дугаарлалт нь стандарт биш маягтад (Монос хүнс) бүлгийн нийт дүн нь гарчиг мөр дээрээ байдаг
HEADER_TOTALS = {norm("Богино хугацаат өр төлбөр"): "total_stl", norm("Урт хугацаат өр төлбөр"): "total_ltl"}
#: Өмч нь эдгээр хэсгээр задарсан маягтад (Монос хүнсний шинэ маягт, "Өмч" мөргүй хуучин маягт) нийлбэр нь хувьцаат капитал.
#: ("Хувьцаат өр төлбөр" нь стандарт маягт дахь "Хувьцаат өмч"-ийн алдаатай нэр.)
SHARE_PARTS = frozenset(map(norm, ("Өмч - Төрийн", "Өмч - Хувийн", "Өмч - Хувьцаат",
                                   "Төрийн өмч", "Хувийн өмч", "Хувьцаат өр төлбөр")))
NUMBER_RE = re.compile(r"\(?-?\d+(\.\d+)?([eE][+-]?\d+)?\)?")
DEFAULT_SCALE = 1000  # нэгж нь тодорхойгүй бол: мянган төгрөг (дараа нь тасралтгүй байдлаар шалгана)

_FILE_PERIOD_RES = (
    re.compile(r"(20\d{2})Q([1-4])", re.I),
    re.compile(r"(20\d{2})([1-4])(?:report)?$", re.I),
    re.compile(r"(20\d{2})0([1-4])$"),   # 'SUU 135202601' = 2026 Q1
)


# ----------------------------------------------------------------------------- үе, огноо
def quarter_end(period: Period) -> date:
    month = period.quarter * 3
    return date(period.year, month, 30 if month in (6, 9) else 31)


def period_of_date(d: date) -> Period:
    return Period(d.year, (d.month - 1) // 3 + 1)


def period_from_filename(path: Path) -> Period:
    stem = re.sub(r"\s+", "", path.stem)
    for rx in _FILE_PERIOD_RES:
        m = rx.search(stem)
        if m:
            return Period(int(m.group(1)), int(m.group(2)))
    return Period(0, 0, path.stem)


def detect_end_date(rows: list[list[Any]]) -> Optional[date]:
    """'Тайлант үе: 2025.01.01-2025.12.31' мэтээс дуусах огноог олно. Мөрд хэд хэдэн хугацаа байвал (өмнөх жил,
    тайлант жил) хамгийн сүүлийнхийг авна."""
    for row in rows[:20]:
        ends: list[date] = []
        for cell in row:
            if not isinstance(cell, str):
                continue
            for rx in (DATE_RANGE_RE, DATE_RANGE_RE2):
                for m in rx.finditer(cell):
                    try:
                        ends.append(date(int(m.group(4)), int(m.group(5)), int(m.group(6))))
                    except ValueError:
                        pass
        if ends:
            return max(ends)
    for row in rows[:20]:                                   # хугацааны мужгүй бол дан огноо ("2026 оны 6-р сарын 30")
        singles: list[date] = []
        for cell in row:
            if isinstance(cell, str):
                for m in DATE_LONG_RE.finditer(cell):
                    try:
                        singles.append(date(int(m.group(1)), int(m.group(2)), int(m.group(3))))
                    except ValueError:
                        pass
        if singles:
            return max(singles)
    return None


def detect_scale(rows: list[list[Any]]) -> Optional[int]:
    """Нэгж -> төгрөгт хөрвүүлэх үржвэр (мянган=1000, сая=1e6, төгрөг=1). Олдохгүй бол None."""
    for row in rows[:20]:
        text = " ".join(c for c in row if isinstance(c, str)).lower()
        if "мян" in text and "төгрөг" in text:
            return 1_000
        if "сая" in text and "төгрөг" in text:
            return 1_000_000
        if "тэрбум" in text and "төгрөг" in text:
            return 1_000_000_000
        if re.search(r"/\s*(төгрөг|tugrik|tugrug|mnt)", text):
            return 1
    return None


# ----------------------------------------------------------------------------- файл унших
def read_sheets(path: Path) -> dict[str, list[list[Any]]]:
    """{sheet_name: мөрүүд}. Форматыг өргөтгөлөөр биш агуулгаар (ZIP = xlsx) тодорхойлно:
    сайт дээр `.xls` нэртэй боловч үнэндээ `.xlsx` файл байдаг."""
    with open(path, "rb") as fh:
        magic = fh.read(4)
    if magic == b"%PDF":                       # текст давхаргатай PDF-г Excel-тэй адил sheet болгоно (pdf.py)
        from .pdf import read_pdf_sheets

        return read_pdf_sheets(path)
    is_xlsx = magic == b"PK\x03\x04"
    if is_xlsx:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(path.read_bytes()), data_only=True, read_only=True)   # өргөтгөлийг openpyxl шалгадаг тул байтаар нь өгнө
        try:
            return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
        finally:
            wb.close()
    import xlrd

    book = xlrd.open_workbook(str(path))
    return {s.name: [s.row_values(i) for i in range(s.nrows)] for s in book.sheets()}


def to_number(value: Any) -> Optional[float]:
    """Тоо эсвэл тоо мэт текст ('92974324.0', '1,234', '(500)') -> float; бусад -> None."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        s = value.replace(",", "").replace(" ", "").replace(" ", "")
        if NUMBER_RE.fullmatch(s):
            neg = s.startswith(("(", "-"))
            val = float(s.strip("()-"))
            return -val if neg else val
    return None


def _code(cell: Any) -> str:
    if cell in (None, ""):
        return ""
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return str(cell).strip()


def find_table(rows: list[list[Any]]) -> Optional[tuple[int, int, Optional[int], int]]:
    """(толгой мөр, үзүүлэлтийн багана, мөрийн дугаарын багана, тайлант утгын багана)."""
    for i, row in enumerate(rows):
        label_col = next((j for j, c in enumerate(row) if isinstance(c, str) and LABEL_HEADER_RE.match(norm(c))), None)
        if label_col is None:
            continue
        code_col = next((j for j, c in enumerate(row) if isinstance(c, str) and CODE_HEADER_RE.match(c.strip().lower())), None)
        filled = [j for j, c in enumerate(row)
                  if c not in (None, "") and not (isinstance(c, str) and NOTE_HEADER_RE.match(norm(c)))]
        if len(filled) >= 3:
            return i, label_col, code_col, filled[-1]  # сүүлийн багана = эцсийн үлдэгдэл / тайлант үеийн дүн
    return None


def find_headerless_table(rows: list[list[Any]]) -> Optional[tuple[int, int, None, int]]:
    """"Үзүүлэлт" толгойгүй маягт (Монос хүнс): нэрний багана хоосон, дараагийн 2+ нүдэнд он/хугацаа бичигдсэн мөр.
    Үзүүлэлт нь A баганад ("1.1.1 Мөнгө..."), утга нь хугацааны баганууд; сүүлийнх нь тайлант үе."""
    for i, row in enumerate(rows[:20]):
        if row and row[0] not in (None, ""):
            continue
        cols = [j for j, c in enumerate(row) if j > 0 and isinstance(c, str) and re.search(r"\d{4}", c)]
        if len(cols) >= 2:
            return i, 0, None, cols[-1]
    return None


def _parts(row: list[Any], label_col: int, code_col: Optional[int]) -> tuple[str, str]:
    """(нэр, мөрийн дугаар). Дугаар нь нэрний нүдэнд нийлсэн бол салгана."""
    label = str(row[label_col]).strip() if row[label_col] not in (None, "") else ""
    code = _code(row[code_col]) if code_col is not None else ""
    if code_col is None and label:
        if m := EMBEDDED_CODE_RE.match(label):
            code, label = m.group(1), m.group(2).strip()
        else:
            label = TOTAL_NUM_RE.sub("Нийт ", label)
    return label, code


def extract_statement(rows: list[list[Any]], statement: str) -> tuple[dict[str, float], list[str]]:
    """Загварын түлхүүр -> түүхий утга (файлд заасан нэгжээр) ба загварт таараагүй утгатай мөрүүд."""
    table = find_table(rows) or find_headerless_table(rows)
    if table is None:
        return {}, []
    head, label_col, code_col, val_col = table
    body = [list(r) + [None] * (val_col + 1 - len(r)) for r in rows[head + 1:]]
    parts = [_parts(r, label_col, code_col) for r in body]
    # Дугаарлалт стандарт эсэх: 2.1.1.x = богино хугацаат өр. Монос хүнсний шинэ маягтад 2.1.2... нь бүгд богино
    # хугацаатын дэд мөр тул дугаараар биш гарчиг мөрөөр бүлгийг тогтооно.
    standard = code_col is not None or any(c.startswith("2.1.1.") for _, c in parts)
    values: dict[str, float] = {}
    unmatched: list[str] = []
    group: Optional[str] = None
    after_oci = False
    share_sum: Optional[float] = None
    for row, (label, code) in zip(body, parts):
        if not label and not code:
            continue
        cur = to_number(row[val_col])
        n = norm(label)
        header_group = GROUP_HEADERS.get(n)
        if standard:
            group = group_of(code) or (header_group if cur in (None, 0.0) else None) or group
        else:
            group = header_group or group
        item = None
        if header_group and cur in (None, 0.0):
            continue                                   # бүлгийн гарчиг мөр (жишээ нь "2.3 Эздийн өмч") — өгөгдөл биш
        if is_english(label):
            after_oci = after_oci or n == norm("Other comprehensive income")
            key = english_key(statement, label, after_oci and n != norm("Other comprehensive income"))
            item = BY_KEY.get(key) if key else None
        else:
            if statement == BS and n in SHARE_PARTS:
                if cur is not None:
                    share_sum = (share_sum or 0.0) + cur
                continue
            item = match_item(statement, label, code, group)
            if item is None and statement == BS and cur not in (None, 0.0) and n in HEADER_TOTALS:
                item = BY_KEY[HEADER_TOTALS[n]]
        if item is None:
            if cur not in (None, 0.0) and label and n not in IGNORED:
                unmatched.append(label)
            continue
        if cur is not None and item.key not in values:
            values[item.key] = cur
    if share_sum is not None and "share_capital" not in values:
        values["share_capital"] = share_sum
    return values, unmatched


# ----------------------------------------------------------------------------- тайлан
def parse_report(item: Downloaded) -> Optional[ParsedReport]:
    """Нэг Excel тайланг стандарт загварт задална. Баланс/ОДТ олдохгүй бол None."""
    sheets = read_sheets(item.path)
    raw: dict[str, tuple[dict[str, float], list[str], Optional[int], Optional[date]]] = {}
    for name, rows in sheets.items():
        name = name.strip()
        if BALANCE_SHEET_RE.search(name) and BS not in raw:
            kind = BS
        elif INCOME_SHEET_RE.search(name) and IS not in raw:
            kind = IS
        else:
            continue
        vals, unmatched = extract_statement(rows, kind)
        if vals:
            raw[kind] = (vals, unmatched, detect_scale(rows), detect_end_date(rows))
    if not raw:
        return None

    file_end = next((r[3] for r in raw.values() if r[3]), None)
    if item.period and item.period.year:
        # Сайтын хүснэгтийн (Он | Улирал) үе нь эх сурвалж: --since шүүлт, давхардал бүгд үүгээр ажиллана.
        period, end = item.period, quarter_end(item.period)
        if file_end and period_of_date(file_end) != period:
            log.warning("%s: файл доторх үе (%s) нь сайтын хүснэгтийн үе (%s)-тэй ялгаатай — хүснэгтийнхийг авлаа",
                        item.path.name, period_of_date(file_end).label, period.label)
    elif file_end:
        # Хүснэгтийн үе байхгүй (гараар хадгалсан файл): файл доторх огноо > файлын нэр
        period = period_of_date(file_end)
        end = quarter_end(period)                  # багана нь улирлын сүүлийн өдрөөр тогтвортой (--since шүүлттэй нийцнэ)
        by_name = period_from_filename(item.path)
        if by_name.year and by_name != period:
            log.warning("%s: файл доторх үе (%s) нь файлын нэрнээс (%s) ялгаатай — файл доторхыг авлаа",
                        item.path.name, period.label, by_name.label)
    else:
        by_name = period_from_filename(item.path)
        if not by_name.year:
            log.warning("%s: тайлангийн үе тодорхойгүй — алгаслаа", item.path.name)
            return None
        period, end = by_name, quarter_end(by_name)

    fallback_scale = next((r[2] for r in raw.values() if r[2]), None) or guess_scale(raw)
    report = ParsedReport(item.ticker, period, end, period.quarter * 3, source=item.path.name, filed=item.filed)
    for kind, (vals, unmatched, scale, _end) in raw.items():
        factor = (scale or fallback_scale or DEFAULT_SCALE) / 1000          # -> мянган төгрөг
        for key, v in vals.items():
            report.values[key] = v if _is_per_share(key) else v * factor
        report.unmatched += unmatched
    v = report.values
    if v.get("net_income") == 0 and v.get("profit_after_tax"):
        # Зарим компани "Тайлант үеийн цэвэр ашиг" мөрийг 0 үлдээдэг (SUU 2024 Q4) — татварын дараах ашгаар орлуулна
        v["net_income"] = v["profit_after_tax"] + (v.get("discontinued") or 0.0)
        report.scale_note = "Цэвэр ашиг 0 байсан → татварын дараах ашгаар орлуулсан"
        log.warning("%s: %s", item.path.name, report.scale_note)
    return report


def guess_scale(raw: dict) -> int:
    """Нэгж бичигдээгүй тайланд: нийт хөрөнгө 5 тэрбумаас их бол бүтэн төгрөг, үгүй бол мянган төгрөг гэж үзнэ.
    (Мянган төгрөгөөр хамгийн том нь ~1.5 тэрбум (APU), бүтэн төгрөгөөр хамгийн жижиг нь ~10 тэрбум.)"""
    assets = raw.get(BS, ({},))[0].get("total_assets")
    return 1 if assets and abs(assets) > 5e9 else DEFAULT_SCALE


def _is_per_share(key: str) -> bool:
    return next(i.kind for i in ITEMS if i.key == key) == "per_share"


def scan_local(reports_dir: Path, tickers: list[str]) -> list[Downloaded]:
    """Өмнө татсан/гараар хадгалсан файлуудыг (татахгүйгээр) олно."""
    items: list[Downloaded] = []
    for path in sorted(reports_dir.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in (".xls", ".xlsx", ".pdf"):
            continue
        folder = path.parent.name.upper()
        first = re.match(r"[A-Za-z]+", path.stem)
        ticker = folder if folder in tickers else (first.group(0).upper() if first else "")
        if ticker in tickers:
            items.append(Downloaded(ticker, path, "", None))
    return items
