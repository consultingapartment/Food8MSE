import logging
from datetime import date
from pathlib import Path

import pytest

from mse_financials.models import Downloaded, Period
from mse_financials.parse import (detect_end_date, detect_scale, extract_statement, find_table, guess_scale,
                                  parse_report, period_from_filename, quarter_end, scan_local, to_number)


@pytest.mark.parametrize("name,expected", [
    ("APU 9020252report.xls", Period(2025, 2)),
    ("MFC  551 20244report.xls", Period(2024, 4)),
    ("TUM 54920254.xlsx", Period(2025, 4)),
    ("MMX 20820142report.xls", Period(2014, 2)),
    ("GAZR 565_20254.xls", Period(2025, 4)),
    ("SUU 135202601.xls", Period(2026, 1)),
    ("APU_2026Q2.xlsx", Period(2026, 2)),
])
def test_period_from_filename(name, expected):
    assert period_from_filename(Path(name)) == expected


def test_unknown_filename_period():
    assert period_from_filename(Path("MFC 551_finance_report_xls_69e5cc83d5a5f.xls")).year == 0


@pytest.mark.parametrize("value,expected", [
    (12, 12.0), ("92974324.0", 92974324.0), ("1,234.5", 1234.5), ("(500)", -500.0), ("-457461.0", -457461.0),
    ("Нийт хөрөнгө", None), ("1.1.1", None), (None, None), (True, None), ("", None),
])
def test_to_number(value, expected):
    assert to_number(value) == expected


def test_quarter_end():
    assert [quarter_end(Period(2025, q)) for q in (1, 2, 3, 4)] == [
        date(2025, 3, 31), date(2025, 6, 30), date(2025, 9, 30), date(2025, 12, 31)]


@pytest.mark.parametrize("cell,expected", [
    ("2025.01.01-2025.12.31", date(2025, 12, 31)),
    ("2025/01/01 - 2025/06/30", date(2025, 6, 30)),
    ("2025.01.01- 2025.09.30", date(2025, 9, 30)),
])
def test_detect_end_date(cell, expected):
    assert detect_end_date([["", "Тайлант үе:", cell]]) == expected
    assert detect_end_date([["нэр"]]) is None


@pytest.mark.parametrize("text,expected", [
    ("/Мянган төгрөг/", 1000), ("/ мянган төгрөг/", 1000), ("/мян.төгрөгөөр/", 1000),
    ("/төгрөгөөр/", 1), ("/ төгрөг/", 1), ("/сая төгрөг/", 1_000_000), ("нэгж алга", None),
])
def test_detect_scale(text, expected):
    assert detect_scale([["", text]]) == expected


def test_guess_scale_by_magnitude():
    assert guess_scale({"balance": ({"total_assets": 1.2e9}, [], None, None)}) == 1000
    assert guess_scale({"balance": ({"total_assets": 6.5e10}, [], None, None)}) == 1


def test_find_table_both_header_styles():
    old = [["нэр"], ["", "№", "Үзүүлэлт", "Эхний үлдэгдэл", "Эцсийн үлдэгдэл"]]
    new = [["Мөрийн дугаар", "Балансын зүйл", "", "Эхний үлдэгдэл", "Эцсийн үлдэгдэл"]]
    assert find_table(old) == (1, 2, 1, 4)
    assert find_table(new) == (0, 1, 0, 4)
    assert find_table([["юмгүй"]]) is None


def test_parse_report_scales_to_thousand_mnt_and_matches_template(tmp_path, report_factory):
    f = report_factory(tmp_path / "APU_2025Q2.xlsx", 1000, 300)
    rep = parse_report(Downloaded("APU", f, "", Period(2025, 2)))
    assert rep.period == Period(2025, 2) and rep.end_date == date(2025, 6, 30) and rep.months == 6
    v = rep.values
    assert v["total_assets"] == 1000 and v["total_liabilities"] + v["total_equity"] == v["total_le"] == 1000
    assert v["revenue"] == 300 and v["net_income"] == 30 and v["cogs"] == 180
    assert v["st_provision"] == 30 and v["lt_provision"] == 10          # нэг нэртэй хоёр мөр ялгагдсан
    assert v["tax_receivable"] == 10 and v["share_capital"] == 100
    assert rep.unmatched == []                                          # -хувийн, 'Бусад дэлгэрэнгүй орлого' = хаясан


def test_full_tugrug_header_is_divided_by_1000(tmp_path, report_factory):
    f = report_factory(tmp_path / "x.xlsx", 5_000_000, 1_000_000, unit="/төгрөгөөр/", scale=1000)
    rep = parse_report(Downloaded("TUM", f, "", Period(2025, 2)))
    assert rep.values["total_assets"] == pytest.approx(5_000_000)       # 5 тэрбум төгрөг = 5 сая мянган төгрөг


def test_file_period_beats_misleading_filename(tmp_path, report_factory, caplog):
    """МХБ-д 'MMX 20820264' нэртэй файл дотроо 2025 оны жилийн тайлан байдаг (бодит тохиолдол)."""
    f = report_factory(tmp_path / "MMX 20820264.xls".replace(".xls", ".xlsx"), 100, 50, end="2025/01/01 - 2025/12/31")
    with caplog.at_level(logging.WARNING):
        rep = parse_report(Downloaded("MMX", f, "", None))                # гараар хадгалсан файл: хүснэгтийн үе байхгүй
    assert rep.period == Period(2025, 4) and rep.end_date == date(2025, 12, 31)
    assert "ялгаатай" in caplog.text


def test_site_table_period_beats_text_inside_file(tmp_path, report_factory, caplog):
    """Сайтын хүснэгтэд 2026 Q2 гэж ирүүлсэн файлын доторх огноо хуучин (хуулбарласан) бол хүснэгтийг баримтална."""
    f = report_factory(tmp_path / "x.xlsx", 100, 50, end="2025/01/01 - 2025/06/30")
    with caplog.at_level(logging.WARNING):
        rep = parse_report(Downloaded("APU", f, "", Period(2026, 2)))
    assert rep.period == Period(2026, 2) and rep.end_date == date(2026, 6, 30)
    assert "хүснэгтийн үе" in caplog.text


def test_no_period_anywhere_is_skipped(tmp_path, report_factory):
    f = report_factory(tmp_path / "MFC 551_finance_report_xls_abc.xlsx", 100, 50, end=None)
    assert parse_report(Downloaded("MFC", f, "", None)) is None


def test_missing_unit_is_guessed_from_magnitude(tmp_path, report_factory):
    f = report_factory(tmp_path / "t.xlsx", 99_000_000_000, 5e10, unit="", scale=1)
    rep = parse_report(Downloaded("TUM", f, "", Period(2025, 4)))
    assert rep.values["total_assets"] == pytest.approx(99_000_000)


def test_english_sheet_names(tmp_path, report_factory):
    f = report_factory(tmp_path / "e.xlsx", 100, 50, sheets=("BS", "IS"))
    assert parse_report(Downloaded("APU", f, "", Period(2025, 2))).values["total_assets"] == 100


def test_numeric_cells_work_too(tmp_path, report_factory):
    f = report_factory(tmp_path / "n.xlsx", 100, 50, string_numbers=False)
    assert parse_report(Downloaded("APU", f, "", Period(2025, 2))).values["revenue"] == 50


def test_extract_statement_returns_unmatched_with_value_only(tmp_path):
    rows = [["№", "Үзүүлэлт", "Эхний", "Эцсийн"], ["1", "Борлуулалтын орлого (цэвэр)", 0, 5],
            ["99", "Үл мэдэгдэх мөр", 0, 7], ["98", "Үл мэдэгдэх тэг мөр", 0, 0]]
    vals, unmatched = extract_statement(rows, "income")
    assert vals == {"revenue": 5.0} and unmatched == ["Үл мэдэгдэх мөр"]


def test_parse_legacy_xls_with_xlrd(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    bs = wb.add_sheet("СБД")
    for i, row in enumerate([["", "/Мянган төгрөг/"], ["", "№", "Үзүүлэлт", "Эхний", "Эцсийн"],
                             ["", "1.3", "Нийт хөрөнгийн дүн", "100.0", "250.0"]]):
        for j, v in enumerate(row):
            bs.write(i, j, v)
    inc = wb.add_sheet("ОДТ")
    for i, row in enumerate([["", "№", "Үзүүлэлт", "Эхний", "Эцсийн"], ["", "1", "Борлуулалтын орлого (цэвэр)", "40.0", "90.0"]]):
        for j, v in enumerate(row):
            inc.write(i, j, v)
    path = tmp_path / "APU 9020252report.xls"
    wb.save(str(path))
    rep = parse_report(Downloaded("APU", path, "", None))
    assert rep.period == Period(2025, 2) and rep.values == {"total_assets": 250.0, "revenue": 90.0}


def test_scan_local_finds_ticker_from_folder_or_filename(tmp_path, report_factory):
    (tmp_path / "SUU").mkdir()
    report_factory(tmp_path / "SUU" / "anything.xlsx", 1, 2)
    report_factory(tmp_path / "APU 9020252report.xlsx", 1, 2)
    report_factory(tmp_path / "XXX 1.xlsx", 1, 2)
    (tmp_path / "manifest.json").write_text("{}")
    found = {(d.ticker, d.path.name) for d in scan_local(tmp_path, ["APU", "SUU"])}
    assert found == {("SUU", "anything.xlsx"), ("APU", "APU 9020252report.xlsx")}


def test_xlsx_content_with_xls_extension_is_read(tmp_path, report_factory):
    """Сайт дээр `.xls` нэртэй боловч агуулга нь xlsx байх тохиолдол (APU 2015 Q4)."""
    real = report_factory(tmp_path / "real.xlsx", 1000, 300)
    disguised = tmp_path / "APU 9020154report.xls"
    disguised.write_bytes(real.read_bytes())
    rep = parse_report(Downloaded("APU", disguised, "", Period(2015, 4)))
    assert rep is not None and rep.values["total_assets"] == 1000


def test_note_column_after_values_is_not_taken_as_value():
    rows = [["", "Үзүүлэлт", "Эхний үлдэгдэл", "Эцсийн үлдэгдэл", "Тайлбар"],
            ["1", "Мөнгө, түүнтэй адилтгах хөрөнгө", "10", "20", "хавсралт 5"]]
    assert extract_statement(rows, "balance")[0]["cash"] == 20.0


def test_report_date_is_normalised_to_quarter_end(tmp_path, report_factory):
    f = report_factory(tmp_path / "x.xlsx", 100, 50, end="2025.01.01-2025.05.31")
    rep = parse_report(Downloaded("APU", f, "", None))
    assert rep.period == Period(2025, 2) and rep.end_date == date(2025, 6, 30)


def test_zero_net_income_row_falls_back_to_profit_after_tax(tmp_path, report_factory):
    f = report_factory(tmp_path / "x.xlsx", 1000, 300)
    rep = parse_report(Downloaded("SUU", f, "", Period(2025, 2)))
    assert rep.values["net_income"] == 30                              # fixture: цэвэр ашиг = 10%
    # ОДТ-ийн "цэвэр ашиг" мөрийг 0 болгоод (бодит SUU 2024 Q4) дахин задална
    from openpyxl import load_workbook
    wb = load_workbook(f)
    ws = wb["ОДТ"]
    for row in ws.iter_rows():
        if row[2].value and "цэвэр ашиг" in str(row[2].value).lower() and "татварын" not in str(row[2].value).lower():
            row[-1].value = "0.0"
    g = tmp_path / "y.xlsx"
    wb.save(g)
    rep2 = parse_report(Downloaded("SUU", g, "", Period(2025, 2)))
    assert "орлуулсан" in rep2.scale_note
