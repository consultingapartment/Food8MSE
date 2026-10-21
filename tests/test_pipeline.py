import logging
import re
from datetime import date

import pytest
from openpyxl import load_workbook

from mse_financials.client import FetchError
from mse_financials.models import ParsedReport, Period
from mse_financials.pipeline import fix_scales, run

H = "https://issuers.mse.mn//uploads/finance_reports"
DAY = date(2026, 10, 10)


class FakeResponse:
    def __init__(self, content: bytes):
        self.content = content
        self.text = content.decode("utf-8", "ignore")


class FakeClient:
    """Сүлжээгүйгээр сайтыг дуурайна."""

    base_url = "https://open.mse.mn"

    def __init__(self, pages: dict[str, bytes]):
        self.pages, self.calls = pages, []

    def get(self, url):
        key = re.sub(r"^https://[^/]+", "", url).replace("//", "/")
        self.calls.append(key)
        for k, v in self.pages.items():
            if key == k.replace("//", "/"):
                return FakeResponse(v)
        raise FetchError(f"{url}: HTTP 404")

    def text(self, url):
        return self.get(url).text

    def sleep(self):
        pass


@pytest.fixture
def site(tmp_path, report_factory):
    """APU: 3 Excel (2 үе + hash нэртэй), 1 PDF; SUU: 1 Excel."""
    def blob(name, assets, revenue, **kw):
        return report_factory(tmp_path / name, assets, revenue, **kw).read_bytes()

    securities = ('<div id="stocks"><div id="class1"><table>'
                  '<tr><td>1</td><td><a href="/securities/90">APU</a></td><td>APU</td></tr>'
                  '<tr><td>2</td><td><a href="/securities/135">SUU</a></td><td>SUU</td></tr></table></div></div>')
    apu = (f'<table><tr><td>2025</td><td>4</td><td>2026-02-28</td><td><a href="{H}/90_finance_report_xls_a1.xlsx">x</a></td></tr>'
           f'<tr><td>2025</td><td>2</td><td>d</td><td><a href="{H}/ 90 20252report.xlsx">x</a></td></tr>'
           f'<tr><td>2024</td><td>4</td><td>d</td><td><a href="{H}/9020244report.xlsx">x</a></td></tr>'
           f'<tr><td>2019</td><td>4</td><td>d</td><td><a href="{H}/9020194report.pdf">pdf</a></td></tr></table>')
    suu = f'<table><tr><td>2025</td><td>4</td><td>2026-03-01</td><td><a href="{H}/13520254report.xlsx">x</a></td></tr></table>'
    return {
        "/securities": securities.encode(),
        "/securities/90/tab/financials": apu.encode(),
        "/securities/135/tab/financials": suu.encode(),
        "/uploads/finance_reports/90_finance_report_xls_a1.xlsx": blob("a.xlsx", 2000, 900, end="2025.01.01-2025.12.31"),
        "/uploads/finance_reports/ 90 20252report.xlsx": blob("b.xlsx", 1500, 400, end=None),
        "/uploads/finance_reports/9020244report.xlsx": blob("c.xlsx", 1200, 800, end="2024.01.01-2024.12.31"),
        "/uploads/finance_reports/9020194report.pdf": b"%PDF-1.4",
        "/uploads/finance_reports/13520254report.xlsx": blob("d.xlsx", 700, 300, end="2025.01.01-2025.12.31"),
    }


def book(path):
    return load_workbook(path)


def header(ws, row):
    return [c.value for c in ws[row]][1:]


def test_one_file_one_sheet_per_company_named_by_company(tmp_path, site):
    out = run(["APU", "SUU", "TUM"], tmp_path / "rep", tmp_path / "reports" / "FOOD mse.xlsx",
              day=DAY, client=FakeClient(site))
    wb = book(out)
    assert wb.sheetnames == ["АПУ ХК", "Түмэн шувуут ХК", "Сүү ХК"]      # config дахь дараалал
    apu, suu = wb["АПУ ХК"], wb["Сүү ХК"]
    assert header(apu, 5) == ["2024 Q4", "2025 Q2", "2025 Q4"]                     # тайлан бүр нэг багана
    assert [d.date() for d in header(apu, 4)] == [date(2024, 12, 31), date(2025, 6, 30), date(2025, 12, 31)]
    assert header(apu, 6) == [12, 6, 12]
    assert [c.number_format for c in apu[4][1:4]] == ["yyyy-mm-dd"] * 3
    assert apu["A7"].value == "Ирүүлсэн огноо (МХБ)"
    assert [c.date() if c else None for c in header(apu, 7)] == [None, None, date(2026, 2, 28)]
    assert header(suu, 5) == ["2025 Q4"]
    assert [c.date() for c in header(suu, 7)] == [date(2026, 3, 1)]
    assert "олдсонгүй" in wb["Түмэн шувуут ХК"]["A4"].value                      # тайлангүй компани


def test_every_sheet_uses_the_same_rows(tmp_path, site):
    wb = book(run(["APU", "SUU"], tmp_path / "rep", tmp_path / "FOOD mse.xlsx", day=DAY, client=FakeClient(site)))
    apu, suu = wb["АПУ ХК"], wb["Сүү ХК"]
    def labels(ws):
        return {ws.cell(row=r, column=1).value: r for r in range(10, ws.max_row + 1) if ws.cell(row=r, column=1).value}
    la, ls = labels(apu), labels(suu)
    assert la == ls                                                              # ижил мөр, ижил дараалал
    assert [apu.cell(row=la["Нийт хөрөнгийн дүн"], column=c).value for c in range(2, 5)] == [1200, 1500, 2000]
    assert [apu.cell(row=la["Борлуулалтын орлого (цэвэр)"], column=c).value for c in range(2, 5)] == [800, 400, 900]
    assert suu.cell(row=ls["Нийт хөрөнгийн дүн"], column=2).value == 700
    assert "Нөөц /өр төлбөр/ (богино хугацаат)" in la and "Нөөц /өр төлбөр/ (урт хугацаат)" in la


def test_balance_check_row_is_formula_and_zero(tmp_path, site):
    ws = book(run(["APU"], tmp_path / "rep", tmp_path / "FOOD mse.xlsx", day=DAY, client=FakeClient(site)))["АПУ ХК"]
    row = next(r for r in range(10, ws.max_row + 1) if str(ws.cell(row=r, column=1).value).startswith("Балансын зөрүү"))
    formula = ws.cell(row=row, column=2).value
    assert formula.startswith("=IF(AND(ISNUMBER(")


def test_rerun_overwrites_the_same_single_file(tmp_path, site):
    out = tmp_path / "reports" / "FOOD mse.xlsx"
    run(["APU"], tmp_path / "rep", out, day=DAY, client=FakeClient(site))
    c2 = FakeClient(site)
    run(["APU"], tmp_path / "rep", out, day=date(2026, 11, 13), client=c2)
    assert [p.name for p in out.parent.iterdir()] == ["FOOD mse.xlsx"]
    assert not [c for c in c2.calls if "uploads" in c]         # манифест: дахин татаагүй


def test_pdf_is_downloaded_but_not_parsed_and_no_pdf_flag(tmp_path, site):
    c = FakeClient(site)
    run(["APU"], tmp_path / "rep", tmp_path / "o.xlsx", day=DAY, client=c)
    assert any(p.suffix == ".pdf" for p in (tmp_path / "rep" / "APU").iterdir())
    c2 = FakeClient(site)
    run(["APU"], tmp_path / "rep2", tmp_path / "o2.xlsx", day=DAY, include_pdf=False, client=c2)
    assert not any(x.endswith(".pdf") for x in c2.calls)


def test_skip_download_parses_local_folder(tmp_path, report_factory):
    folder = tmp_path / "rep"
    folder.mkdir()
    report_factory(folder / "APU 9020252report.xlsx", 1000, 300)
    report_factory(folder / "APU 9020254.xlsx", 1100, 700, end="2025.01.01-2025.12.31")
    ws = book(run(["APU"], folder, tmp_path / "o.xlsx", day=DAY, skip_download=True))["АПУ ХК"]
    assert header(ws, 5) == ["2025 Q2", "2025 Q4"]


def test_corrupt_file_does_not_stop_run(tmp_path, report_factory):
    folder = tmp_path / "rep"
    folder.mkdir()
    (folder / "APU 9020252report.xlsx").write_bytes(b"not an excel file")
    report_factory(folder / "APU 9020254.xlsx", 10, 20, end="2025.01.01-2025.12.31")
    ws = book(run(["APU"], folder, tmp_path / "o.xlsx", day=DAY, skip_download=True))["АПУ ХК"]
    assert header(ws, 5) == ["2025 Q4"]


def test_duplicate_period_keeps_most_complete(tmp_path, report_factory):
    folder = tmp_path / "rep"
    folder.mkdir()
    report_factory(folder / "APU 9020254.xlsx", 10, 20, end="2025.01.01-2025.12.31")
    report_factory(folder / "APU 90202541.xlsx", 99, 20, end="2025.01.01-2025.12.31", sheets=("BS", "IS"))
    ws = book(run(["APU"], folder, tmp_path / "o.xlsx", day=DAY, skip_download=True))["АПУ ХК"]
    assert len(header(ws, 5)) == 1


def mk(ticker, y, assets):
    return ParsedReport(ticker, Period(y, 4), date(y, 12, 31), 12,
                        {"total_assets": float(assets), "revenue": assets / 2})


def test_fix_scales_detects_stale_unit_header():
    reps = [mk("MFC", 2022, 40_000), mk("MFC", 2023, 44_000), mk("MFC", 2024, 47_000_000), mk("MFC", 2025, 50_000)]
    fixed = fix_scales(reps)
    assert [r.period.year for r in fixed] == [2024]
    assert reps[2].values["total_assets"] == pytest.approx(47_000) and "÷1000" in reps[2].scale_note
    assert reps[2].values["revenue"] == pytest.approx(23_500)
    assert reps[0].scale_note == ""


def test_fix_scales_treats_balance_and_income_statement_separately():
    """SUU 2024 Q4: баланс бүтэн төгрөгөөр (1000x), харин ОДТ зөв нэгжтэй — зөвхөн балансыг засна."""
    reps = [mk("SUU", y, a) for y, a in ((2022, 40_000), (2023, 44_000), (2025, 50_000))]
    odd = ParsedReport("SUU", Period(2024, 4), date(2024, 12, 31), 12,
                       {"total_assets": 47_000_000.0, "total_equity": 20_000_000.0, "revenue": 23_500.0})
    fix_scales(reps + [odd])
    assert odd.values["total_assets"] == pytest.approx(47_000) and odd.values["total_equity"] == pytest.approx(20_000)
    assert odd.values["revenue"] == pytest.approx(23_500)                      # ОДТ хөндөгдөөгүй
    assert odd.scale_note.startswith("Баланс ÷1000") and "ОДТ" not in odd.scale_note


def test_fix_scales_fixes_income_statement_by_turnover():
    """Баланс зөв, харин ОДТ 1000 дахин том: орлого/хөрөнгө = 1000x → зөвхөн ОДТ-ийг ÷1000."""
    reps = [mk("TUM", y, a) for y, a in ((2022, 40_000), (2023, 44_000), (2025, 50_000))]
    odd = ParsedReport("TUM", Period(2024, 4), date(2024, 12, 31), 12,
                       {"total_assets": 47_000.0, "revenue": 30_000_000.0, "net_income": 3_000_000.0})
    fix_scales(reps + [odd])
    assert odd.values["total_assets"] == 47_000 and odd.values["revenue"] == pytest.approx(30_000)
    assert odd.values["net_income"] == pytest.approx(3_000) and odd.scale_note.startswith("ОДТ ÷1000")


def test_fix_scales_leaves_genuine_growth_alone():
    reps = [mk("X", 2021, 10_000), mk("X", 2022, 25_000), mk("X", 2023, 60_000), mk("X", 2024, 140_000)]
    assert fix_scales(reps) == []


def test_overdue_mandatory_report_is_warned(tmp_path, site, caplog):
    with caplog.at_level(logging.WARNING):
        run(["APU"], tmp_path / "rep", tmp_path / "o.xlsx", day=DAY, client=FakeClient(site))
    assert "APU: 2026 Q2" in caplog.text                       # сайтад 2026 Q2 алга
    caplog.clear()
    with caplog.at_level(logging.WARNING):
        run(["APU"], tmp_path / "rep", tmp_path / "o.xlsx", day=date(2026, 2, 1), client=FakeClient(site))
    assert "2026 Q2" not in caplog.text


def test_identical_figures_in_two_periods_are_flagged(tmp_path, report_factory, caplog):
    folder = tmp_path / "rep"
    folder.mkdir()
    report_factory(folder / "APU 9020252report.xlsx", 1000, 300)
    report_factory(folder / "APU 9020254report.xlsx", 1000, 300, end="2025.01.01-2025.12.31")      # хуучин тайланг буруу хавсаргасан мэт
    with caplog.at_level("WARNING"):
        run(["APU"], folder, tmp_path / "o.xlsx", day=DAY, skip_download=True)
    assert any("бүрэн ижил" in m for m in caplog.messages)


def _labels(ws):
    return {ws.cell(row=r, column=1).value: r for r in range(10, ws.max_row + 1) if ws.cell(row=r, column=1).value}


def test_incremental_update_appends_only_the_new_report(tmp_path, report_factory):
    folder, out = tmp_path / "rep", tmp_path / "reports" / "FOOD mse.xlsx"
    folder.mkdir()
    report_factory(folder / "APU 9020252report.xlsx", 1000, 300)
    run(["APU"], folder, out, day=DAY, skip_download=True)
    assert header(book(out)["АПУ ХК"], 5) == ["2025 Q2"]

    report_factory(folder / "APU 9020254report.xlsx", 2000, 900, end="2025.01.01-2025.12.31")
    run(["APU"], folder, out, day=DAY, skip_download=True)
    ws = book(out)["АПУ ХК"]
    assert header(ws, 5) == ["2025 Q2", "2025 Q4"]
    row = _labels(ws)["Нийт хөрөнгийн дүн"]
    assert [ws.cell(row=row, column=c).value for c in (2, 3)] == [1000, 2000]      # хуучин багана хэвээр


def test_nothing_new_leaves_file_untouched_and_logs_it(tmp_path, report_factory, caplog):
    folder, out = tmp_path / "rep", tmp_path / "FOOD mse.xlsx"
    folder.mkdir()
    report_factory(folder / "APU 9020252report.xlsx", 1000, 300)
    run(["APU", "SUU"], folder, out, day=DAY, skip_download=True)
    before = out.read_bytes()
    with caplog.at_level("INFO"):
        run(["APU", "SUU"], folder, out, day=DAY, skip_download=True)
    assert out.read_bytes() == before
    assert any("нэгтгэгдсэн — шинэ тайлан алга" in m for m in caplog.messages)      # APU
    assert any(m.startswith("SUU: sheet хоосон") for m in caplog.messages)


def test_online_rerun_skips_already_merged_periods(tmp_path, site):
    out = tmp_path / "reports" / "FOOD mse.xlsx"
    run(["APU"], tmp_path / "rep", out, day=DAY, client=FakeClient(site))
    c2 = FakeClient(site)
    run(["APU"], tmp_path / "rep2", out, day=date(2026, 11, 13), client=c2)          # хоосон кэш, гэхдээ Excel байна
    assert not [c for c in c2.calls if "uploads" in c]                               # шинэ үе байхгүй -> юу ч татаагүй
    assert header(book(out)["АПУ ХК"], 5) == ["2024 Q4", "2025 Q2", "2025 Q4"]


def test_rebuild_ignores_existing_workbook(tmp_path, report_factory):
    folder, out = tmp_path / "rep", tmp_path / "FOOD mse.xlsx"
    folder.mkdir()
    report_factory(folder / "APU 9020252report.xlsx", 1000, 300)
    run(["APU"], folder, out, day=DAY, skip_download=True)
    (folder / "APU 9020252report.xlsx").unlink()
    report_factory(folder / "APU 9020254report.xlsx", 2000, 900, end="2025.01.01-2025.12.31")
    run(["APU"], folder, out, day=DAY, skip_download=True, rebuild=True)
    assert header(book(out)["АПУ ХК"], 5) == ["2025 Q4"]


def test_partial_run_keeps_other_companies_sheets(tmp_path, report_factory):
    folder, out = tmp_path / "rep", tmp_path / "FOOD mse.xlsx"
    folder.mkdir()
    report_factory(folder / "APU 9020252report.xlsx", 1000, 300)
    report_factory(folder / "SUU 13520252report.xlsx", 500, 100)
    run(["APU", "SUU"], folder, out, day=DAY, skip_download=True)
    report_factory(folder / "APU 9020254report.xlsx", 2000, 900, end="2025.01.01-2025.12.31")
    run(["APU"], folder, out, day=DAY, skip_download=True)                # зөвхөн APU
    wb = book(out)
    assert wb.sheetnames == ["АПУ ХК", "Сүү ХК"]
    assert header(wb["АПУ ХК"], 5) == ["2025 Q2", "2025 Q4"] and header(wb["Сүү ХК"], 5) == ["2025 Q2"]


def test_older_gap_period_is_not_inserted_behind_the_last_column(tmp_path, report_factory):
    folder, out = tmp_path / "rep", tmp_path / "FOOD mse.xlsx"
    folder.mkdir()
    report_factory(folder / "APU 9020254report.xlsx", 2000, 900, end="2025.01.01-2025.12.31")
    run(["APU"], folder, out, day=DAY, skip_download=True)
    report_factory(folder / "APU 9020252report.xlsx", 1000, 300)          # хуучин үе — сүүлийн баганаас өмнөх
    run(["APU"], folder, out, day=DAY, skip_download=True)
    assert header(book(out)["АПУ ХК"], 5) == ["2025 Q4"]                  # --rebuild үед л нөхнө


def test_inconsistent_subtotals_are_flagged(tmp_path, report_factory, caplog):
    from mse_financials.pipeline import warn_inconsistent
    f = report_factory(tmp_path / "APU 9020252report.xlsx", 1000, 300)
    from mse_financials.models import Downloaded
    from mse_financials.parse import parse_report
    rep = parse_report(Downloaded("APU", f, "", None))
    with caplog.at_level("WARNING"):
        warn_inconsistent([rep])
    assert not caplog.messages                                   # зөв тайланд сануулга байхгүй
    rep.values["total_equity"] += 500
    with caplog.at_level("WARNING"):
        warn_inconsistent([rep])
    assert any("Хөрөнгө ≠ Өр + Эздийн өмч" in m for m in caplog.messages)



def test_failed_tickers_are_reported_to_caller(tmp_path, site):
    failed: list[str] = []
    run(["APU", "TUM"], tmp_path / "rep", tmp_path / "o.xlsx", day=DAY, client=FakeClient(site), failed=failed)
    assert failed == ["TUM"]                                      # TUM хуудас 404 — CLI exit code 1 болгоно


def test_top_level_statement_headings_are_kept(tmp_path, report_factory):
    folder = tmp_path / "rep"
    folder.mkdir()
    report_factory(folder / "APU 9020252report.xlsx", 1000, 300)
    ws = book(run(["APU"], folder, tmp_path / "o.xlsx", day=DAY, skip_download=True))["АПУ ХК"]
    labels = [ws.cell(row=r, column=1).value for r in range(10, ws.max_row + 1)]
    assert labels[0].startswith("САНХҮҮГИЙН БАЙДЛЫН ТАЙЛАН") and labels[1] == "Эргэлтийн хөрөнгө"
    assert any(str(x).startswith("ОРЛОГЫН ДЭЛГЭРЭНГҮЙ ТАЙЛАН") for x in labels)


# ── --check, summary, CLI ────────────────────────────────────────────────────

def site_with_new_q2(site, tmp_path, report_factory):
    """APU-ийн хуудсанд 2026 Q2 нэмсэн хувилбар."""
    s = dict(site)
    row = f'<tr><td>2026</td><td>2</td><td>2026-07-28</td><td><a href="{H}/9020262report.xlsx">x</a></td></tr>'
    s["/securities/90/tab/financials"] = site["/securities/90/tab/financials"].decode().replace("<table>", "<table>" + row, 1).encode()
    s["/uploads/finance_reports/9020262report.xlsx"] = report_factory(
        tmp_path / "n.xlsx", 2200, 500, end="2026.01.01-2026.06.30").read_bytes()
    return s


def test_check_lists_only_newer_reports_and_downloads_nothing(tmp_path, site, report_factory):
    from mse_financials.pipeline import check
    out = tmp_path / "reports" / "FOOD mse.xlsx"
    run(["APU"], tmp_path / "rep", out, day=DAY, client=FakeClient(site))
    before = out.read_bytes()
    s2 = site_with_new_q2(site, tmp_path, report_factory)
    c = FakeClient(s2)
    res = check(["APU"], out, client=c)
    assert res["APU"]["new"] == [("2026 Q2", {"xls"})]
    assert res["APU"]["last"] == date(2025, 12, 31) and res["APU"]["error"] is None
    assert not [x for x in c.calls if "uploads" in x]          # файл татаагүй
    assert out.read_bytes() == before                           # Excel хөндөгдөөгүй


def test_check_without_workbook_reports_everything_and_flags_pdf_only(tmp_path, site):
    from mse_financials.pipeline import check
    res = check(["APU"], tmp_path / "missing.xlsx", client=FakeClient(site))
    assert [lbl for lbl, _ in res["APU"]["new"]] == ["2019 Q4", "2024 Q4", "2025 Q2", "2025 Q4"]
    assert dict(res["APU"]["new"])["2019 Q4"] == {"pdf"}
    assert "2019 Q4" not in [lbl for lbl, _ in check(["APU"], tmp_path / "missing.xlsx",
                                                       client=FakeClient(site), include_pdf=False)["APU"]["new"]]


def test_check_reports_fetch_errors_per_ticker(tmp_path, site):
    from mse_financials.pipeline import check
    s = {k: v for k, v in site.items() if k != "/securities/135/tab/financials"}
    res = check(["APU", "SUU"], tmp_path / "missing.xlsx", client=FakeClient(s))
    assert res["APU"]["error"] is None and "404" in res["SUU"]["error"]


def test_run_fills_summary(tmp_path, site, report_factory):
    out = tmp_path / "FOOD mse.xlsx"
    s1: dict = {}
    run(["APU"], tmp_path / "rep", out, day=DAY, client=FakeClient(site), summary=s1)
    assert s1["added"] == {"APU": ["2024 Q4", "2025 Q2", "2025 Q4"]}
    s2: dict = {}
    run(["APU"], tmp_path / "rep", out, day=DAY, client=FakeClient(site), summary=s2)
    assert s2["added"] == {} and s2["last"]["APU"] == date(2025, 12, 31)
    s3: dict = {}
    run(["APU"], tmp_path / "rep", out, day=DAY, client=FakeClient(site_with_new_q2(site, tmp_path, report_factory)),
        summary=s3)
    assert s3["added"] == {"APU": ["2026 Q2"]} and s3["last"]["APU"] == date(2026, 6, 30)


def test_cli_check_does_not_touch_excel_and_notifies(tmp_path, site, monkeypatch, capsys):
    from mse_financials import cli
    out = tmp_path / "FOOD mse.xlsx"
    run(["APU"], tmp_path / "rep", out, day=DAY, client=FakeClient(site))
    before = out.read_bytes()
    monkeypatch.setattr(cli, "Client", lambda **kw: FakeClient(site))
    sent = []
    monkeypatch.setattr(cli, "send_telegram", lambda t, **kw: sent.append(t) or True)
    rc = cli.main(["--check", "--tickers", "APU", "--out", str(out), "--notify", "--date", "2026-10-12"])
    assert rc == 0 and out.read_bytes() == before
    assert "APU" in capsys.readouterr().out and len(sent) == 1 and "APU" in sent[0]


def test_cli_check_with_skip_download_is_rejected(tmp_path, capsys):
    from mse_financials import cli
    assert cli.main(["--check", "--skip-download", "--out", str(tmp_path / "x.xlsx")]) == 2


def test_cli_check_exit_code_1_on_error(tmp_path, site, monkeypatch):
    from mse_financials import cli
    s = {k: v for k, v in site.items() if k != "/securities/90/tab/financials"}
    monkeypatch.setattr(cli, "Client", lambda **kw: FakeClient(s))
    assert cli.main(["--check", "--tickers", "APU", "--out", str(tmp_path / "x.xlsx")]) == 1


def test_cli_run_sends_summary_with_warnings_and_passes_strict_ssl(tmp_path, site, monkeypatch):
    from mse_financials import cli
    seen = {}
    monkeypatch.setattr(cli, "Client", lambda **kw: seen.update(kw) or FakeClient(site))
    sent = []
    monkeypatch.setattr(cli, "send_telegram", lambda t, **kw: sent.append(t) or True)
    rc = cli.main(["--tickers", "APU", "--out", str(tmp_path / "FOOD mse.xlsx"), "--reports-dir", str(tmp_path / "rep"),
                   "--notify", "--strict-ssl", "--date", "2026-10-12"])
    assert rc == 0 and seen == {"strict_ssl": True}
    assert len(sent) == 1 and "3 шинэ тайлан" in sent[0] and "APU" in sent[0]


def test_cli_run_failure_is_notified_and_exit_1(tmp_path, site, monkeypatch):
    from mse_financials import cli
    s = {k: v for k, v in site.items() if k != "/securities/90/tab/financials"}
    monkeypatch.setattr(cli, "Client", lambda **kw: FakeClient(s))
    sent = []
    monkeypatch.setattr(cli, "send_telegram", lambda t, **kw: sent.append(t) or True)
    rc = cli.main(["--tickers", "APU", "--out", str(tmp_path / "o.xlsx"), "--reports-dir", str(tmp_path / "rep"),
                   "--notify", "--date", "2026-10-12"])
    assert rc == 1 and "❌ APU" in sent[0]


def test_pdf_is_not_downloaded_when_excel_exists_for_same_period(tmp_path, site):
    site["/securities/135/tab/financials"] = (
        f'<table><tr><td>2025</td><td>4</td><td>2026-03-01</td><td><a href="{H}/13520254report.pdf">p</a></td>'
        f'<td><a href="{H}/13520254report.xlsx">x</a></td></tr></table>').encode()
    site["/uploads/finance_reports/13520254report.pdf"] = b"%PDF-1.4"
    c = FakeClient(site)
    run(["SUU"], tmp_path / "rep", tmp_path / "o.xlsx", day=DAY, client=c)
    assert any(x.endswith("13520254report.xlsx") for x in c.calls)
    assert not any(x.endswith("13520254report.pdf") for x in c.calls)       # Excel байгаа тул PDF/OCR хэрэггүй


def test_pdf_is_downloaded_when_period_has_only_pdf_or_excel_is_really_a_pdf(tmp_path, site):
    site["/securities/135/tab/financials"] = (
        f'<table><tr><td>2025</td><td>4</td><td>d</td><td><a href="{H}/13520254report.pdf">p</a></td></tr>'
        f'<tr><td>2025</td><td>2</td><td>d</td><td><a href="{H}/13520252report.pdf">p</a></td>'
        f'<td><a href="{H}/13520252report.xls">x</a></td></tr></table>').encode()
    site["/uploads/finance_reports/13520254report.pdf"] = b"%PDF-1.4"
    site["/uploads/finance_reports/13520252report.pdf"] = b"%PDF-1.4 real"
    site["/uploads/finance_reports/13520252report.xls"] = b"%PDF-1.4 disguised"           # ".xls" нэртэй боловч PDF
    c = FakeClient(site)
    run(["SUU"], tmp_path / "rep", tmp_path / "o.xlsx", day=DAY, client=c)
    assert any(x.endswith("13520254report.pdf") for x in c.calls)               # зөвхөн PDF-тэй үе → татна
    assert any(x.endswith("13520252report.pdf") for x in c.calls)               # Excel нь PDF болсон → PDF-ийг ч татна
