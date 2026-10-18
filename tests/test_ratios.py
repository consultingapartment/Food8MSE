"""Харьцааны шинжилгээний блок: томьёоны бүтэц + (LibreOffice байвал) тооцоолсон утга."""
import shutil
import subprocess
from datetime import date

import pytest
from openpyxl import load_workbook

from mse_financials.pipeline import run
from mse_financials.ratios import GROUPS
from mse_financials.template import ITEMS

DAY = date(2026, 3, 1)


def make_two_years(tmp_path, report_factory):
    folder = tmp_path / "rep"
    folder.mkdir()
    report_factory(folder / "APU 9020244report.xlsx", 1000, 800, end="2024.01.01-2024.12.31")
    report_factory(folder / "APU 9020254report.xlsx", 2000, 900, end="2025.01.01-2025.12.31")
    return run(["APU"], folder, tmp_path / "FOOD mse.xlsx", day=DAY, skip_download=True)


def test_ratio_labels_do_not_collide_with_statement_rows():
    """Эс бусадаас: харьцааны нэр нь загварын мөрийн нэртэй давхцвал дахин унших үед утга болж ойлгогдоно."""
    labels = {i.label for i in ITEMS}
    assert not any(spec[1] in labels for _, specs in GROUPS for spec in specs)


def test_ratio_block_has_formulas(tmp_path, report_factory):
    ws = load_workbook(make_two_years(tmp_path, report_factory))["АПУ ХК"]
    rows = {ws.cell(row=r, column=1).value: r for r in range(10, ws.max_row + 1) if ws.cell(row=r, column=1).value}
    gm = next(r for k, r in rows.items() if str(k).startswith("Нийт ашгийн маржин"))
    f = ws.cell(row=gm, column=3).value
    assert f.startswith("=IF(AND(ISNUMBER(") and "IFERROR" in f
    assert any(str(k).startswith("1. Ашигт ажиллагаа") for k in rows)
    assert any(str(k).startswith("6. Өсөлт") for k in rows)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice байхгүй")
def test_ratio_values_after_recalculation(tmp_path, report_factory):
    src = make_two_years(tmp_path, report_factory)
    out = tmp_path / "calc"
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src)],
                   check=True, capture_output=True, timeout=180)
    ws = load_workbook(out / src.name, data_only=True)["АПУ ХК"]
    rows = {str(ws.cell(row=r, column=1).value).split(" =")[0]: r for r in range(10, ws.max_row + 1) if ws.cell(row=r, column=1).value}
    col25 = 3                                                    # 2025 Q4
    val = lambda name: ws.cell(row=rows[name], column=col25).value
    assert val("Нийт ашгийн маржин") == pytest.approx(0.4)                      # fixture: өртөг = 60%
    assert val("Дундаж нийт хөрөнгө, мянган төгрөг") == pytest.approx((1000 + 2000) / 2)     # өмнөх оны Q4 ба эцсийн
    assert val("Өр ÷ Нийт хөрөнгө") == pytest.approx(0.4)
    assert val("Орлогын өсөлт") == pytest.approx(900 / 800 - 1)
    assert val("Жилийн коэффициент (12 ÷ хамрах сар)") == pytest.approx(1.0)
